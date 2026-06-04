import sys
import mysql.connector
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QFrame, 
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView, 
    QDateEdit, QMessageBox, QApplication, QTabWidget, QFileDialog,
    QComboBox, QDialog, QFormLayout, QDoubleSpinBox,
    QDialogButtonBox, QAbstractItemView
)
from PyQt6.QtCore import Qt, QDate, QTimer, QDateTime, QTime
from PyQt6.QtGui import QColor
from PyQt6.QtWidgets import QGraphicsDropShadowEffect
from decimal import Decimal, ROUND_HALF_UP

try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


class SalaryAllowanceDialog(QDialog):
    def __init__(self, db_config, parent=None):
        super().__init__(parent)
        self.db_config = db_config
        self.setWindowTitle("Quản lý lương & phụ cấp nhân viên")
        self.resize(800, 500)
        layout = QVBoxLayout(self)
        self.table = QTableWidget(0, 5)
        self.table.setHorizontalHeaderLabels(["ID NV", "Tên nhân viên", "Lương giờ (đ)", "Phụ cấp (đ)", "Hành động"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)
        btn_add = QPushButton("➕ Thêm nhân viên mới")
        btn_add.clicked.connect(self.add_employee)
        layout.addWidget(btn_add)
        self.load_data()

    def load_data(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("""
                SELECT u.id, u.ho_ten, 
                       COALESCE(ln.luong_theo_gio, 25000) as luong_theo_gio,
                       COALESCE(ln.phu_cap, 0) as phu_cap
                FROM user u
                LEFT JOIN luong_nhanvien ln ON u.id = ln.id_nhanvien
                WHERE u.vai_tro IN ('NhanVien', 'Admin', 'Quản trị')
                ORDER BY u.id
            """)
            rows = cursor.fetchall()
            self.table.setRowCount(len(rows))
            for i, row in enumerate(rows):
                self.table.setItem(i, 0, QTableWidgetItem(str(row['id'])))
                self.table.setItem(i, 1, QTableWidgetItem(row['ho_ten']))
                self.table.setItem(i, 2, QTableWidgetItem(f"{row['luong_theo_gio']:,.0f}"))
                self.table.setItem(i, 3, QTableWidgetItem(f"{row['phu_cap']:,.0f}"))
                btn_widget = QWidget()
                btn_layout = QHBoxLayout(btn_widget)
                btn_layout.setContentsMargins(5, 5, 5, 5)
                btn_edit = QPushButton("Sửa")
                btn_edit.setStyleSheet("background-color: #3498db; color: white;")
                btn_edit.clicked.connect(lambda checked, uid=row['id']: self.edit_employee(uid))
                btn_delete = QPushButton("Xóa")
                btn_delete.setStyleSheet("background-color: #e74c3c; color: white;")
                btn_delete.clicked.connect(lambda checked, uid=row['id']: self.delete_employee(uid))
                btn_layout.addWidget(btn_edit)
                btn_layout.addWidget(btn_delete)
                self.table.setCellWidget(i, 4, btn_widget)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))
        finally:
            db.close()

    def connect_db(self):
        try:
            return mysql.connector.connect(**self.db_config)
        except mysql.connector.Error as err:
            QMessageBox.critical(self, "Lỗi kết nối DB", str(err))
            return None

    def add_employee(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("""
                SELECT id, ho_ten FROM user 
                WHERE vai_tro IN ('NhanVien', 'Admin', 'Quản trị')
                AND id NOT IN (SELECT id_nhanvien FROM luong_nhanvien)
            """)
            employees = cursor.fetchall()
            if not employees:
                QMessageBox.information(self, "Thông báo", "Tất cả nhân viên đã có thông tin lương.")
                return
            dialog = QDialog(self)
            dialog.setWindowTitle("Chọn nhân viên để thêm lương")
            layout = QVBoxLayout(dialog)
            combo = QComboBox()
            for emp in employees:
                combo.addItem(f"{emp['id']} - {emp['ho_ten']}", emp['id'])
            layout.addWidget(QLabel("Chọn nhân viên:"))
            layout.addWidget(combo)
            form = QFormLayout()
            spin_luong = QDoubleSpinBox()
            spin_luong.setRange(0, 1000000)
            spin_luong.setValue(25000)
            spin_luong.setSuffix(" đ")
            spin_phucap = QDoubleSpinBox()
            spin_phucap.setRange(0, 1000000)
            spin_phucap.setValue(0)
            spin_phucap.setSuffix(" đ")
            form.addRow("Lương theo giờ:", spin_luong)
            form.addRow("Phụ cấp:", spin_phucap)
            layout.addLayout(form)
            buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
            buttons.accepted.connect(dialog.accept)
            buttons.rejected.connect(dialog.reject)
            layout.addWidget(buttons)
            if dialog.exec() == QDialog.DialogCode.Accepted:
                nv_id = combo.currentData()
                luong = spin_luong.value()
                phucap = spin_phucap.value()
                cursor.execute("INSERT INTO luong_nhanvien (id_nhanvien, luong_theo_gio, phu_cap) VALUES (%s, %s, %s)",
                               (nv_id, luong, phucap))
                db.commit()
                QMessageBox.information(self, "Thành công", "Đã thêm thông tin lương cho nhân viên.")
                self.load_data()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))
        finally:
            db.close()

    def edit_employee(self, nv_id):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("SELECT luong_theo_gio, phu_cap FROM luong_nhanvien WHERE id_nhanvien = %s", (nv_id,))
            row = cursor.fetchone()
            if not row:
                QMessageBox.warning(self, "Lỗi", "Không tìm thấy thông tin lương.")
                return
            dialog = QDialog(self)
            dialog.setWindowTitle(f"Sửa lương & phụ cấp - ID {nv_id}")
            layout = QVBoxLayout(dialog)
            form = QFormLayout()
            spin_luong = QDoubleSpinBox()
            spin_luong.setRange(0, 1000000)
            spin_luong.setValue(row['luong_theo_gio'])
            spin_luong.setSuffix(" đ")
            spin_phucap = QDoubleSpinBox()
            spin_phucap.setRange(0, 1000000)
            spin_phucap.setValue(row['phu_cap'])
            spin_phucap.setSuffix(" đ")
            form.addRow("Lương theo giờ:", spin_luong)
            form.addRow("Phụ cấp:", spin_phucap)
            layout.addLayout(form)
            buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
            buttons.accepted.connect(dialog.accept)
            buttons.rejected.connect(dialog.reject)
            layout.addWidget(buttons)
            if dialog.exec() == QDialog.DialogCode.Accepted:
                cursor.execute("UPDATE luong_nhanvien SET luong_theo_gio = %s, phu_cap = %s WHERE id_nhanvien = %s",
                               (spin_luong.value(), spin_phucap.value(), nv_id))
                db.commit()
                QMessageBox.information(self, "Thành công", "Đã cập nhật.")
                self.load_data()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))
        finally:
            db.close()

    def delete_employee(self, nv_id):
        reply = QMessageBox.question(self, "Xác nhận", f"Xóa thông tin lương của nhân viên ID {nv_id}?",
                                     QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if reply == QMessageBox.StandardButton.Yes:
            db = self.connect_db()
            if not db:
                return
            try:
                cursor = db.cursor()
                cursor.execute("DELETE FROM luong_nhanvien WHERE id_nhanvien = %s", (nv_id,))
                db.commit()
                QMessageBox.information(self, "Thành công", "Đã xóa.")
                self.load_data()
            except Exception as e:
                QMessageBox.critical(self, "Lỗi", str(e))
            finally:
                db.close()


class AttendanceManagement(QWidget):
    def __init__(self, user_info=None):
        super().__init__()
        self.db_config = {
            "host": "localhost",
            "user": "root",
            "password": "",
            "database": "quanly_snack_db"
        }
        self.user_info = user_info or {"id": 5, "ho_ten": "NGUYỄN THỊ KHANH ", "vai_tro": "Quản trị"}
        self.setWindowTitle("Hệ thống Quản lý Nhân sự & Tài chính - Đồ án của Khanh")
        self.resize(1400, 900)
        self.setStyleSheet("background-color: #f4f4f4; font-family: 'Segoe UI';")
        self.main_layout = QVBoxLayout(self)
        self.main_layout.setContentsMargins(0, 0, 0, 0)
        self.main_layout.setSpacing(0)

        self.header = QFrame()
        self.header.setFixedHeight(70)
        self.header.setStyleSheet("background-color: #58111A; color: white;")
        header_layout = QHBoxLayout(self.header)
        title_lbl = QLabel("HỆ THỐNG QUẢN LÝ LƯƠNG & LỢI NHUẬN")
        title_lbl.setStyleSheet("font-size: 20px; font-weight: bold; padding-left: 20px;")
        user_lbl = QLabel(f"Người dùng: {self.user_info['ho_ten']}  ")
        header_layout.addWidget(title_lbl)
        header_layout.addStretch()
        header_layout.addWidget(user_lbl)
        self.main_layout.addWidget(self.header)

        self.container = QWidget()
        self.container_layout = QVBoxLayout(self.container)
        self.main_layout.addWidget(self.container)

        vai_tro = str(self.user_info.get('vai_tro', '')).strip().lower()
        if any(x in vai_tro for x in ["quản lý", "admin", "quản trị"]):
            self.init_admin_ui()
        else:
            self.init_staff_ui()

    def create_shadow(self, widget):
        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(15)
        shadow.setXOffset(0)
        shadow.setYOffset(4)
        shadow.setColor(QColor(0, 0, 0, 50))
        widget.setGraphicsEffect(shadow)

    def init_admin_ui(self):
        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("""
            QTabWidget::pane { border: none; background: white; }
            QTabBar::tab { 
                background: #e0e0e0; color: #333; padding: 12px 30px; 
                font-weight: bold; border-top-left-radius: 8px; margin-right: 5px;
            }
            QTabBar::tab:selected { background: #58111A; color: white; }
        """)
        self.setup_attendance_tab()
        self.setup_finance_tab()
        self.setup_all_data_tab()
        self.tabs.addTab(self.tab_att, "Quản lý Chấm Công")
        self.tabs.addTab(self.tab_finance, "Tài Chính & Lợi Nhuận")
        self.tabs.addTab(self.tab_all_data, "📊 Toàn bộ dữ liệu chấm công")
        self.container_layout.addWidget(self.tabs)
        self.load_attendance_data()

    def setup_attendance_tab(self):
        self.tab_att = QWidget()
        layout_att = QVBoxLayout(self.tab_att)
        filter_box = QFrame()
        filter_box.setStyleSheet("background: white; border-radius: 10px;")
        self.create_shadow(filter_box)
        filter_layout = QHBoxLayout(filter_box)
        self.date_filter = QDateEdit(QDate.currentDate())
        self.date_filter.setCalendarPopup(True)
        self.date_filter.setFixedSize(150, 35)
        self.date_filter.dateChanged.connect(self.load_attendance_data)
        btn_reload = QPushButton("🔄 Cập nhật")
        btn_reload.clicked.connect(self.load_attendance_data)
        btn_reload.setStyleSheet("background: #58111A; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        btn_excel_att = QPushButton("📎 Xuất Excel")
        btn_excel_att.clicked.connect(lambda: self.export_table_to_excel(self.table, "Bang_ChamCong"))
        btn_excel_att.setStyleSheet("background: #27ae60; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        btn_pdf_att = QPushButton("📄 Xuất PDF")
        btn_pdf_att.clicked.connect(lambda: self.export_table_to_pdf(self.table, "Bang_ChamCong"))
        btn_pdf_att.setStyleSheet("background: #58111A; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        filter_layout.addWidget(QLabel("Ngày:"))
        filter_layout.addWidget(self.date_filter)
        filter_layout.addStretch()
        filter_layout.addWidget(btn_reload)
        filter_layout.addWidget(btn_excel_att)
        filter_layout.addWidget(btn_pdf_att)
        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["ID", "Nhân Viên", "Ngày", "Giờ Vào", "Giờ Ra", "Trạng Thái"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setStyleSheet("QHeaderView::section { background-color: #58111A; color: white; font-weight: bold; }")
        layout_att.addWidget(filter_box)
        layout_att.addWidget(self.table)

    def setup_finance_tab(self):
        self.tab_finance = QWidget()
        layout_fin = QVBoxLayout(self.tab_finance)
        dash_layout = QHBoxLayout()
        self.cards = []
        stats = [("DOANH THU", "0đ", "#1e8449"), ("GIÁ VỐN", "0đ", "#a93226"), 
                 ("TỔNG LƯƠNG", "0đ", "#2e86c1"), ("LỢI NHUẬN RÒNG", "0đ", "#58111A")]
        for title, val, color in stats:
            card = QFrame()
            card.setStyleSheet(f"background: white; border-left: 8px solid {color}; border-radius: 8px;")
            self.create_shadow(card)
            v_lay = QVBoxLayout(card)
            v_lay.addWidget(QLabel(title, styleSheet="color: #7f8c8d; font-size: 11px; font-weight: bold;"))
            v_lbl = QLabel(val)
            v_lbl.setStyleSheet(f"color: {color}; font-size: 18px; font-weight: bold;")
            v_lay.addWidget(v_lbl)
            dash_layout.addWidget(card)
            self.cards.append(v_lbl)

        tool_fin = QHBoxLayout()
        self.month_filter = QDateEdit(QDate.currentDate())
        self.month_filter.setDisplayFormat("MM/yyyy")
        btn_run = QPushButton("⚡ XUẤT BÁO CÁO")
        btn_run.setStyleSheet("background: #58111A; color: white; font-weight: bold; padding: 10px 20px; border-radius: 5px;")
        btn_run.clicked.connect(self.load_finance_data)
        btn_manage_salary = QPushButton("🔧 Quản lý lương & phụ cấp")
        btn_manage_salary.setStyleSheet("background: #f39c12; color: white; font-weight: bold; padding: 10px 20px; border-radius: 5px;")
        btn_manage_salary.clicked.connect(self.open_salary_dialog)
        btn_excel_fin = QPushButton("📎 Xuất Excel Lương")
        btn_excel_fin.clicked.connect(lambda: self.export_table_to_excel(self.salary_table, "Bang_Luong"))
        btn_excel_fin.setStyleSheet("background: #27ae60; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        btn_pdf_fin = QPushButton("📄 Xuất PDF Lương")
        btn_pdf_fin.clicked.connect(lambda: self.export_table_to_pdf(self.salary_table, "Bang_Luong"))
        btn_pdf_fin.setStyleSheet("background: #58111A; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        tool_fin.addWidget(QLabel("Chọn tháng:"))
        tool_fin.addWidget(self.month_filter)
        tool_fin.addWidget(btn_run)
        tool_fin.addWidget(btn_manage_salary)
        tool_fin.addStretch()
        tool_fin.addWidget(btn_excel_fin)
        tool_fin.addWidget(btn_pdf_fin)

        self.salary_table = QTableWidget(0, 6)
        self.salary_table.setHorizontalHeaderLabels(["ID", "Nhân Viên", "Tổng Giờ", "Đơn Giá/h", "Phụ cấp", "Thành Tiền"])
        self.salary_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        layout_fin.addLayout(dash_layout)
        layout_fin.addLayout(tool_fin)
        layout_fin.addWidget(QLabel("<b>CHI TIẾT LƯƠNG (Điều kiện: >1h mới tính lương giờ, phụ cấp vẫn cộng nếu có giờ):</b>"))
        layout_fin.addWidget(self.salary_table)

    def setup_all_data_tab(self):
        self.tab_all_data = QWidget()
        layout = QVBoxLayout(self.tab_all_data)

        # Khung thống kê đẹp ở đầu tab
        stats_frame = QFrame()
        stats_frame.setStyleSheet("background-color: #ffffff; border-radius: 10px; padding: 10px;")
        self.create_shadow(stats_frame)
        stats_layout = QHBoxLayout(stats_frame)
        self.stats_label = QLabel("📊 Đang tải dữ liệu thống kê...")
        self.stats_label.setStyleSheet("font-size: 14px; font-weight: bold; color: #2c3e50;")
        stats_layout.addWidget(self.stats_label)
        stats_layout.addStretch()
        layout.addWidget(stats_frame)

        # Toolbar: nút tải lại + xuất Excel + xuất PDF
        toolbar = QHBoxLayout()
        btn_refresh = QPushButton("🔄 Tải lại dữ liệu chấm công")
        btn_refresh.clicked.connect(self.load_all_attendance_data)
        btn_refresh.setStyleSheet("background: #58111A; color: white; padding: 8px 15px; border-radius: 5px; font-weight: bold;")
        
        btn_excel_all = QPushButton("📎 Xuất Excel toàn bộ")
        btn_excel_all.clicked.connect(self.export_all_att_to_excel)
        btn_excel_all.setStyleSheet("background: #27ae60; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        
        btn_pdf_all = QPushButton("📄 Xuất PDF toàn bộ")
        btn_pdf_all.clicked.connect(self.export_all_att_to_pdf)
        btn_pdf_all.setStyleSheet("background: #58111A; color: white; border-radius: 5px; font-weight: bold; padding: 8px 15px;")
        
        toolbar.addWidget(btn_refresh)
        toolbar.addStretch()
        toolbar.addWidget(btn_excel_all)
        toolbar.addWidget(btn_pdf_all)
        layout.addLayout(toolbar)

        self.all_att_table = QTableWidget(0, 6)
        self.all_att_table.setHorizontalHeaderLabels(["ID", "Nhân Viên", "Ngày", "Giờ Vào", "Giờ Ra", "Trạng Thái"])
        self.all_att_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.all_att_table.setStyleSheet("QHeaderView::section { background-color: #58111A; color: white; font-weight: bold; }")
        layout.addWidget(self.all_att_table)

        self.load_all_attendance_data()
    def load_all_attendance_data(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("SELECT * FROM chamcong ORDER BY ngay_lam_viec DESC, gio_vao DESC")
            rows = cursor.fetchall()
            self.all_att_table.setRowCount(len(rows))
            for i, r in enumerate(rows):
                self.all_att_table.setItem(i, 0, QTableWidgetItem(str(r['id_nhanvien'])))
                self.all_att_table.setItem(i, 1, QTableWidgetItem(r['ten_nhanvien']))
                self.all_att_table.setItem(i, 2, QTableWidgetItem(str(r['ngay_lam_viec'])))
                self.all_att_table.setItem(i, 3, QTableWidgetItem(str(r['gio_vao']).split(' ')[-1] if r['gio_vao'] else ""))
                self.all_att_table.setItem(i, 4, QTableWidgetItem(str(r['gio_ra']).split(' ')[-1] if r['gio_ra'] else "Đang làm"))
                self.all_att_table.setItem(i, 5, QTableWidgetItem(r['trang_thai']))
            self.update_statistics()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi tải dữ liệu", str(e))
        finally:
            db.close()

    def update_statistics(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("SELECT SUM(tong_tien) as total_revenue FROM hoadon")
            total_revenue = cursor.fetchone()['total_revenue'] or 0

            cursor.execute("""
                SELECT id_nhanvien, SUM(TIMESTAMPDIFF(SECOND, gio_vao, gio_ra)) as total_seconds
                FROM chamcong
                WHERE gio_ra IS NOT NULL
                GROUP BY id_nhanvien
            """)
            working_data = cursor.fetchall()

            cursor.execute("SELECT id_nhanvien, luong_theo_gio, phu_cap FROM luong_nhanvien")
            salary_rows = cursor.fetchall()
            salary_dict = {}
            for row in salary_rows:
                salary_dict[row['id_nhanvien']] = (float(row['luong_theo_gio'] or 25000), float(row['phu_cap'] or 0))
            DEFAULT_HOURLY = 25000.0

            total_salary = 0.0
            for w in working_data:
                nv_id = w['id_nhanvien']
                total_sec = w['total_seconds']
                if total_sec is None:
                    total_sec = 0
                elif isinstance(total_sec, Decimal):
                    total_sec = float(total_sec)
                else:
                    total_sec = float(total_sec)
                hours = total_sec / 3600.0
                hourly_rate, allowance = salary_dict.get(nv_id, (DEFAULT_HOURLY, 0.0))
                if hours >= 1.0:
                    total_salary += hours * hourly_rate + allowance
                elif hours > 0:
                    total_salary += allowance

            cursor.execute("SELECT COUNT(*) as cnt FROM user WHERE vai_tro='NhanVien'")
            num_employees = cursor.fetchone()['cnt'] or 0
            cursor.execute("SELECT COUNT(*) as cnt FROM sanpham")
            num_products = cursor.fetchone()['cnt'] or 0
            cursor.execute("SELECT COUNT(*) as cnt FROM donhang")
            num_orders = cursor.fetchone()['cnt'] or 0

            stats_text = (f"📊 TỔNG HỢP: Tổng doanh thu: {total_revenue:,.0f}đ | "
                          f"Tổng lương: {total_salary:,.0f}đ | "
                          f"👥 Nhân viên: {num_employees} | 🍔 Sản phẩm: {num_products} | 🧾 Đơn hàng: {num_orders}")
            self.stats_label.setText(stats_text)
        except Exception as e:
            self.stats_label.setText(f"❌ Lỗi tính thống kê: {str(e)}")
        finally:
            db.close()

    def open_salary_dialog(self):
        dialog = SalaryAllowanceDialog(self.db_config, self)
        dialog.exec()
        self.load_finance_data()

    # ------------------- Các hàm xuất Excel, PDF -------------------
    def get_table_data_as_list(self, table_widget):
        headers = []
        for col in range(table_widget.columnCount()):
            header_item = table_widget.horizontalHeaderItem(col)
            headers.append(header_item.text() if header_item else f"Column {col+1}")
        rows = []
        for row in range(table_widget.rowCount()):
            row_data = []
            for col in range(table_widget.columnCount()):
                item = table_widget.item(row, col)
                row_data.append(item.text() if item else "")
            rows.append(row_data)
        return headers, rows

    def export_table_to_excel(self, table_widget, default_filename):
        if not OPENPYXL_AVAILABLE:
            QMessageBox.warning(self, "Thiếu thư viện", "pip install openpyxl")
            return
        if table_widget.rowCount() == 0:
            QMessageBox.information(self, "Không có dữ liệu", "Bảng trống.")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Lưu file Excel", f"{default_filename}_{QDate.currentDate().toString('yyyyMMdd')}.xlsx",
            "Excel Files (*.xlsx)"
        )
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = default_filename
            headers, rows = self.get_table_data_as_list(table_widget)
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
            wb.save(file_path)
            QMessageBox.information(self, "Thành công", f"Đã xuất:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất Excel", str(e))

    def export_table_to_pdf(self, table_widget, default_filename):
        if not REPORTLAB_AVAILABLE:
            QMessageBox.warning(self, "Thiếu thư viện", "pip install reportlab")
            return
        if table_widget.rowCount() == 0:
            QMessageBox.information(self, "Không có dữ liệu", "Bảng trống.")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Lưu file PDF", f"{default_filename}_{QDate.currentDate().toString('yyyyMMdd')}.pdf",
            "PDF Files (*.pdf)"
        )
        if not file_path:
            return
        try:
            headers, rows = self.get_table_data_as_list(table_widget)
            data = [headers] + rows
            doc = SimpleDocTemplate(file_path, pagesize=landscape(A4))
            elements = []
            styles = getSampleStyleSheet()
            title = Paragraph(f"Báo cáo: {default_filename}", styles['Title'])
            elements.append(title)
            elements.append(Spacer(1, 12))
            table = Table(data)
            style = TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#58111A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ])
            table.setStyle(style)
            elements.append(table)
            doc.build(elements)
            QMessageBox.information(self, "Thành công", f"Đã xuất PDF:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất PDF", str(e))

    def export_all_att_to_excel(self):
        """Xuất toàn bộ bảng chấm công ra Excel"""
        if not OPENPYXL_AVAILABLE:
            QMessageBox.warning(self, "Thiếu thư viện", "pip install openpyxl")
            return
        if self.all_att_table.rowCount() == 0:
            QMessageBox.information(self, "Không có dữ liệu", "Bảng chấm công trống.")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Lưu file Excel", f"ToanBo_ChamCong_{QDate.currentDate().toString('yyyyMMdd')}.xlsx",
            "Excel Files (*.xlsx)"
        )
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "ToanBo_ChamCong"
            headers = []
            for col in range(self.all_att_table.columnCount()):
                header_item = self.all_att_table.horizontalHeaderItem(col)
                headers.append(header_item.text() if header_item else f"Column {col+1}")
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            for row_idx in range(self.all_att_table.rowCount()):
                for col_idx in range(self.all_att_table.columnCount()):
                    item = self.all_att_table.item(row_idx, col_idx)
                    value = item.text() if item else ""
                    ws.cell(row=row_idx+2, column=col_idx+1, value=value)
            wb.save(file_path)
            QMessageBox.information(self, "Thành công", f"Đã xuất toàn bộ chấm công ra:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất Excel", str(e))

    def export_all_att_to_pdf(self):
        """Xuất toàn bộ bảng chấm công ra PDF"""
        if not REPORTLAB_AVAILABLE:
            QMessageBox.warning(self, "Thiếu thư viện", "pip install reportlab")
            return
        if self.all_att_table.rowCount() == 0:
            QMessageBox.information(self, "Không có dữ liệu", "Bảng chấm công trống.")
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Lưu file PDF", f"ToanBo_ChamCong_{QDate.currentDate().toString('yyyyMMdd')}.pdf",
            "PDF Files (*.pdf)"
        )
        if not file_path:
            return
        try:
            headers = []
            for col in range(self.all_att_table.columnCount()):
                header_item = self.all_att_table.horizontalHeaderItem(col)
                headers.append(header_item.text() if header_item else f"Column {col+1}")
            data = [headers]
            for row_idx in range(self.all_att_table.rowCount()):
                row_data = []
                for col_idx in range(self.all_att_table.columnCount()):
                    item = self.all_att_table.item(row_idx, col_idx)
                    row_data.append(item.text() if item else "")
                data.append(row_data)
            doc = SimpleDocTemplate(file_path, pagesize=landscape(A4))
            elements = []
            styles = getSampleStyleSheet()
            title = Paragraph("Báo cáo toàn bộ chấm công", styles['Title'])
            elements.append(title)
            elements.append(Spacer(1, 12))
            table = Table(data)
            style = TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#58111A')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ])
            table.setStyle(style)
            elements.append(table)
            doc.build(elements)
            QMessageBox.information(self, "Thành công", f"Đã xuất PDF toàn bộ chấm công:\n{file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Lỗi xuất PDF", str(e))
    def load_finance_data(self):
        thang = self.month_filter.date().month()
        nam = self.month_filter.date().year()
        db = self.connect_db()
        if not db:
            QMessageBox.critical(self, "Lỗi", "Không thể kết nối CSDL.")
            return
        try:
            cursor = db.cursor(dictionary=True)
            cursor.execute("SELECT id_nhanvien, luong_theo_gio, phu_cap FROM luong_nhanvien")
            luong_rows = cursor.fetchall()
            luong_map = {}
            phu_cap_map = {}
            for row in luong_rows:
                nv_id = row['id_nhanvien']
                luong_map[nv_id] = float(row['luong_theo_gio'] or 0)
                phu_cap_map[nv_id] = float(row['phu_cap'] or 0)
            DEFAULT_HOURLY_RATE = 25000.0

            query_salary = """
                SELECT id_nhanvien, ten_nhanvien, 
                SUM(TIMESTAMPDIFF(SECOND, gio_vao, gio_ra)) as total_seconds
                FROM chamcong 
                WHERE MONTH(ngay_lam_viec)=%s AND YEAR(ngay_lam_viec)=%s 
                AND gio_vao IS NOT NULL AND gio_ra IS NOT NULL
                GROUP BY id_nhanvien, ten_nhanvien
            """
            cursor.execute(query_salary, (thang, nam))
            res_salary = cursor.fetchall()
            total_salary_all = 0.0
            self.salary_table.setRowCount(len(res_salary))

            for i, r in enumerate(res_salary):
                nv_id = r['id_nhanvien']
                total_seconds = r['total_seconds'] or 0
                if isinstance(total_seconds, Decimal):
                    total_seconds = float(total_seconds)
                else:
                    total_seconds = float(total_seconds)
                gio = total_seconds / 3600.0
                gio = round(gio, 2)
                gio_float = gio
                don_gia = luong_map.get(nv_id, DEFAULT_HOURLY_RATE)
                phu_cap = phu_cap_map.get(nv_id, 0.0)

                if gio_float <= 0:
                    thanh_tien = 0.0
                    note = " (Không làm ca nào)"
                elif gio_float < 1.0:
                    thanh_tien = phu_cap
                    note = " (Dưới 1h, chỉ phụ cấp)"
                else:
                    thanh_tien = gio_float * don_gia + phu_cap
                    note = ""

                thanh_tien = round(thanh_tien)
                total_salary_all += thanh_tien

                self.salary_table.setItem(i, 0, QTableWidgetItem(str(nv_id)))
                self.salary_table.setItem(i, 1, QTableWidgetItem(r['ten_nhanvien']))
                self.salary_table.setItem(i, 2, QTableWidgetItem(f"{gio_float:.2f}h{note}"))
                self.salary_table.setItem(i, 3, QTableWidgetItem(f"{don_gia:,.0f}đ"))
                self.salary_table.setItem(i, 4, QTableWidgetItem(f"{phu_cap:,.0f}đ"))
                item_money = QTableWidgetItem(f"{thanh_tien:,.0f}đ")
                if thanh_tien == 0:
                    item_money.setForeground(QColor("red"))
                self.salary_table.setItem(i, 5, item_money)

            cursor.execute("SELECT SUM(tong_tien) as dt FROM hoadon WHERE MONTH(ngay_tao)=%s AND YEAR(ngay_tao)=%s", (thang, nam))
            doanh_thu = float(cursor.fetchone()['dt'] or 0)
            gia_von = doanh_thu * 0.35
            chi_phi_khac = 300000
            loi_nhuan = doanh_thu - gia_von - total_salary_all - chi_phi_khac

            self.cards[0].setText(f"{doanh_thu:,.0f}đ")
            self.cards[1].setText(f"{gia_von:,.0f}đ")
            self.cards[2].setText(f"{total_salary_all:,.0f}đ")
            self.cards[3].setText(f"{loi_nhuan:,.0f}đ")

        except Exception as e:
            QMessageBox.critical(self, "Lỗi xử lý dữ liệu", str(e))
        finally:
            db.close()

    def connect_db(self):
        try:
            return mysql.connector.connect(**self.db_config)
        except mysql.connector.Error as err:
            QMessageBox.critical(self, "Lỗi kết nối DB", f"Không thể kết nối MySQL:\n{err}")
            return None

    def load_attendance_data(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor(dictionary=True)
            ngay = self.date_filter.date().toString("yyyy-MM-dd")
            cursor.execute("SELECT * FROM chamcong WHERE ngay_lam_viec = %s", (ngay,))
            rows = cursor.fetchall()
            self.table.setRowCount(len(rows))
            for i, r in enumerate(rows):
                self.table.setItem(i, 0, QTableWidgetItem(str(r['id_nhanvien'])))
                self.table.setItem(i, 1, QTableWidgetItem(r['ten_nhanvien']))
                self.table.setItem(i, 2, QTableWidgetItem(str(r['ngay_lam_viec'])))
                self.table.setItem(i, 3, QTableWidgetItem(str(r['gio_vao']).split(' ')[-1] if r['gio_vao'] else ""))
                self.table.setItem(i, 4, QTableWidgetItem(str(r['gio_ra']).split(' ')[-1] if r['gio_ra'] else "Đang làm"))
                self.table.setItem(i, 5, QTableWidgetItem(r['trang_thai']))
        except Exception as e:
            QMessageBox.critical(self, "Lỗi tải dữ liệu", str(e))
        finally:
            db.close()

    def init_staff_ui(self):
        v_lay = QVBoxLayout()
        v_lay.setAlignment(Qt.AlignmentFlag.AlignCenter)
        card = QFrame()
        card.setFixedSize(450, 350)
        card.setStyleSheet("background: white; border-radius: 20px;")
        self.create_shadow(card)
        c_lay = QVBoxLayout(card)
        c_lay.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_clock = QLabel("00:00:00")
        self.lbl_clock.setStyleSheet("font-size: 55px; font-weight: bold; color: #58111A;")
        btn_in = QPushButton("📍 VÀO CA")
        btn_out = QPushButton("🚪 KẾT THÚC")
        for b in [btn_in, btn_out]:
            b.setFixedSize(280, 55)
            b.setStyleSheet("background: #58111A; color: white; font-weight: bold; border-radius: 10px;")
        btn_out.setStyleSheet("background: #444; color: white; font-weight: bold; border-radius: 10px;")
        btn_in.clicked.connect(self.handle_check_in)
        btn_out.clicked.connect(self.handle_check_out)
        c_lay.addWidget(QLabel("GIỜ HỆ THỐNG", styleSheet="color: #999;"))
        c_lay.addWidget(self.lbl_clock)
        c_lay.addSpacing(20)
        c_lay.addWidget(btn_in)
        c_lay.addWidget(btn_out)
        v_lay.addWidget(card)
        self.container_layout.addLayout(v_lay)
        self.timer = QTimer(self)
        self.timer.timeout.connect(lambda: self.lbl_clock.setText(QTime.currentTime().toString("HH:mm:ss")))
        self.timer.start(1000)

    def handle_check_in(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor()
            now = QDateTime.currentDateTime()
            cursor.execute("""
                INSERT INTO chamcong (id_nhanvien, ten_nhanvien, ngay_lam_viec, gio_vao, trang_thai) 
                VALUES (%s, %s, %s, %s, %s)
            """, (self.user_info['id'], self.user_info['ho_ten'],
                  now.toString("yyyy-MM-dd"), now.toString("yyyy-MM-dd HH:mm:ss"), "Ổn định"))
            db.commit()
            QMessageBox.information(self, "Thông báo", "Chào mừng bạn vào ca!")
        except mysql.connector.IntegrityError:
            QMessageBox.warning(self, "Lỗi", "Bạn đã vào ca hôm nay rồi.")
        except Exception as e:
            QMessageBox.warning(self, "Lỗi", str(e))
        finally:
            db.close()

    def handle_check_out(self):
        db = self.connect_db()
        if not db:
            return
        try:
            cursor = db.cursor()
            now = QDateTime.currentDateTime().toString("yyyy-MM-dd HH:mm:ss")
            cursor.execute("""
                UPDATE chamcong SET gio_ra = %s 
                WHERE id_nhanvien = %s AND DATE(ngay_lam_viec) = CURDATE() AND gio_ra IS NULL
            """, (now, self.user_info['id']))
            db.commit()
            if cursor.rowcount > 0:
                QMessageBox.information(self, "Thông báo", "Hẹn gặp lại bạn!")
            else:
                QMessageBox.warning(self, "Lỗi", "Không tìm thấy ca làm việc hôm nay để kết thúc.")
        except Exception as e:
            QMessageBox.warning(self, "Lỗi", str(e))
        finally:
            db.close()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = AttendanceManagement(user_info={"id": 5, "ho_ten": "NGUYỄN THỊ KHANH", "vai_tro": "Quản trị"})
    window.show()
    sys.exit(app.exec())